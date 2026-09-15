"""Extract submitted files inside the isolated, read-only document runtime."""

import json
import subprocess
from pathlib import Path

import openpyxl
import pdfplumber
from pptx import Presentation

root = Path("/files")
results = []
for path in sorted(root.rglob("*")):
    if not path.is_file() or path.is_symlink() or not path.resolve().is_relative_to(root):
        continue
    if path.suffix.lower() not in {
        ".docx",
        ".xlsx",
        ".pdf",
        ".pptx",
        ".txt",
        ".md",
        ".csv",
        ".json",
    }:
        continue
    record = {"filename": str(path.relative_to(root)), "bytes": path.stat().st_size}
    try:
        if path.suffix == ".docx":
            p = subprocess.run(
                ["pandoc", str(path), "-t", "markdown", "--wrap=none", "--track-changes=accept"],
                capture_output=True,
                text=True,
                check=True,
                timeout=30,
            )
            record["text"] = p.stdout
        elif path.suffix == ".xlsx":
            wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
            record["sheets"] = {
                ws.title: {"rows": ws.max_row, "columns": ws.max_column, "values": list(ws.values)}
                for ws in wb
            }
            wb.close()
            cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
            record["cached_values"] = {ws.title: list(ws.values) for ws in cached}
            cached.close()
        elif path.suffix == ".pdf":
            with pdfplumber.open(path) as pdf:
                record["pages"] = [page.extract_text() or "" for page in pdf.pages]
        elif path.suffix == ".pptx":
            deck = Presentation(path)
            record["slides"] = [
                [s.text for s in slide.shapes if s.has_text_frame] for slide in deck.slides
            ]
        else:
            record["text"] = path.read_text()
    except Exception as exc:
        record["parse_error"] = str(exc)
    results.append(record)
print(json.dumps(results, default=str, ensure_ascii=False))
